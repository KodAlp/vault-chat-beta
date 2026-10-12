# -*- coding: utf-8 -*-
"""Tests for images.py ("images survive conversion").

    python ingest/images_test.py

Same convention as convert_test.py: not part of `npm test`. Covers alt-text
filtering, the noise filters, the note block format, token <-> image mapping
on a generated PPTX (MarkItDown run for real) and on a synthetic docling
document (the serializer must equal docling's own export apart from the
pictures).
"""
import base64
import io
import random
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import images  # noqa: E402
from images import Candidate  # noqa: E402

failures: list[str] = []


def check(name: str, cond: bool) -> None:
    print(("ok   " if cond else "FAIL ") + name)
    if not cond:
        failures.append(name)


def noise_png(w: int, h: int, seed: int = 0, mode: str = "RGB", fmt: str = "PNG") -> bytes:
    """Incompressible content, so the 2 KB size filter does not interfere."""
    from PIL import Image

    rnd = random.Random(seed)
    img = Image.new(mode, (w, h))
    bands = len(img.getbands())
    img.putdata([tuple(rnd.randrange(256) for _ in range(bands)) if bands > 1 else rnd.randrange(256) for _ in range(w * h)])
    out = io.BytesIO()
    img.save(out, fmt)
    return out.getvalue()


def test_alt_filter() -> None:
    tr = "metin, ekran görüntüsü, logo içeren bir resim Yapay zeka tarafından oluşturulan içerik yanlış olabilir."
    en = "A screenshot of a computer AI-generated content may be incorrect."
    check("auto alt (Turkish suffix) is dropped", images.clean_alt(tr) == "")
    check("auto alt (English suffix) is dropped", images.clean_alt(en) == "")
    check("suffix match is case-insensitive", images.clean_alt("x AI-GENERATED CONTENT MAY BE INCORRECT") == "")
    check("shape-name fallback is dropped", images.clean_alt("Picture 12") == "" and images.clean_alt("Resim4") == "")
    check("file name fallback is dropped", images.clean_alt("image1.png") == "")
    check("'containing' alone is NOT enough to drop", images.clean_alt("A picture containing the Raft state machine") != "")
    check("real alt text survives", images.clean_alt("Şekil: leader seçimi") == "Şekil: leader seçimi")
    check("brackets/newlines are flattened", images.clean_alt("a [b]\nc") == "a b c")
    check("long alt is shortened", len(images.clean_alt("kelime " * 100)) <= images.ALT_MAX_CHARS + 1)


def test_reencode() -> None:
    from PIL import Image

    blob, ext = images.reencode(noise_png(80, 80))
    check("png stays png", ext == "png" and blob[:4] == b"\x89PNG")
    blob, ext = images.reencode(noise_png(80, 80, fmt="JPEG"))
    check("jpeg stays jpeg", ext == "jpg" and blob[:3] == b"\xff\xd8\xff")
    cmyk = Image.frombytes("CMYK", (300, 300), random.Random(5).randbytes(300 * 300 * 4))
    buf = io.BytesIO()
    cmyk.save(buf, "JPEG")
    blob, ext = images.reencode(buf.getvalue())
    check("CMYK jpeg becomes RGB jpeg", ext == "jpg" and Image.open(io.BytesIO(blob)).mode == "RGB")
    pal = Image.frombytes("P", (80, 80), random.Random(1).randbytes(6400))
    pal.putpalette(list(random.Random(2).randbytes(768)))
    buf = io.BytesIO()
    pal.save(buf, "PNG", transparency=0)
    blob, ext = images.reencode(buf.getvalue())
    check("palette png with transparency becomes RGBA png", ext == "png" and Image.open(io.BytesIO(blob)).mode == "RGBA")
    # metadata is dropped
    img = Image.open(io.BytesIO(noise_png(80, 80, fmt="JPEG")))
    exif = Image.Exif()
    exif[0x010E] = "secret description"
    buf = io.BytesIO()
    img.save(buf, "JPEG", exif=exif)
    blob, _ = images.reencode(buf.getvalue())
    check("EXIF is stripped", b"secret description" not in blob and not Image.open(io.BytesIO(blob)).getexif())
    for name, data, reason in [
        ("tiny side", noise_png(47, 200), "tiny"),
        ("empty data", b"", "unreadable"),
        ("not an image", b"%PDF-1.4 not an image" * 200, "unreadable"),
        ("truncated png", noise_png(120, 120)[:300], "unreadable"),
    ]:
        try:
            images.reencode(data)
            check(f"reencode rejects {name}", False)
        except images.SkipImage as skip:
            check(f"reencode rejects {name} ({reason})", skip.reason == reason)
    # a uniform 100x100 PNG is far below 2 KB: size filter
    flat = io.BytesIO()
    Image.new("RGB", (100, 100), (3, 4, 5)).save(flat, "PNG")
    try:
        images.reencode(flat.getvalue())
        check("small file (<2KB) is rejected", False)
    except images.SkipImage as skip:
        check("small file (<2KB) is rejected", skip.reason == "tiny")
    # decompression bomb protection: over twice the pixel ceiling raises inside PIL
    old = images.MAX_PIXELS
    images.MAX_PIXELS = 1000
    try:
        images.reencode(noise_png(200, 200))
        check("pixel ceiling skips the image", False)
    except images.SkipImage as skip:
        check("pixel ceiling skips the image", skip.reason == "unreadable")
    finally:
        images.MAX_PIXELS = old


def test_equation_images() -> None:
    from PIL import Image

    blob, ext = images.reencode(noise_png(300, 30, seed=7))
    check("300x30 equation-like image is kept", ext == "png" and bool(blob))
    flat = io.BytesIO()
    Image.new("RGB", (200, 24), (255, 255, 255)).save(flat, "PNG")
    check("small-bytes wide equation strip is kept", len(flat.getvalue()) < images.MIN_BYTES and bool(images.reencode(flat.getvalue())[0]))
    for name, data in [
        ("30x30 icon", noise_png(30, 30, seed=8)),
        ("tiny wide 100x6 rule", noise_png(100, 6, seed=9)),
        ("small near-square 100x47", noise_png(100, 47, seed=10)),
    ]:
        try:
            images.reencode(data)
            check(f"{name} is dropped", False)
        except images.SkipImage as skip:
            check(f"{name} is dropped", skip.reason == "tiny")


def test_select_filters() -> None:
    logo = noise_png(100, 100, seed=1)
    frame = noise_png(120, 90, seed=2)
    cands = [Candidate(logo, p, "", area_frac=0.05) for p in (1, 2, 3, 4)]  # small + 4 pages -> template
    cands += [Candidate(frame, p, "", area_frac=0.9) for p in (1, 2, 3, 4)]  # huge + 4 pages -> content
    cands += [Candidate(noise_png(100, 100, seed=3), p, "", area_frac=0.05) for p in (5, 6)]  # 2 pages -> content
    cands += [Candidate(noise_png(30, 30, seed=4), 7, "")]  # tiny
    cands += [Candidate(b"junk", 7, "")]  # unreadable
    kinds = [d.kind for d in images.select(cands, total_pages=8)]
    check("small image on >=3 pages is template", kinds[:4] == ["template"] * 4)
    check("full-slide image repeated is content", kinds[4:8] == ["kept"] * 4)
    check("image on 2 pages is kept at both", kinds[8:10] == ["kept", "kept"])
    check("tiny and unreadable", kinds[10:] == ["tiny", "unreadable"])
    unknown = [Candidate(logo, p, "") for p in (1, 2, 3)]  # area unknown -> never template
    check("unknown area never counts as template", all(d.kind == "kept" for d in images.select(unknown, total_pages=3)))
    small = [Candidate(logo, p, "", area_frac=0.05) for p in (1, 2, 3)]
    check("small figure on 3 of 50 pages is content, not template", all(d.kind == "kept" for d in images.select(small, total_pages=50)))
    check("unknown page count: nothing is template", all(d.kind == "kept" for d in images.select(small)))
    capped = images.select([Candidate(noise_png(60, 60, seed=s), 1, "") for s in range(6)], max_kept=4)
    check("cap keeps the first N", [d.kind for d in capped] == ["kept"] * 4 + ["over_cap"] * 2)


def test_format_block() -> None:
    sha = "ab" * 8
    check("block with description", images.format_block(sha, "png", "alt", "alt") == f"![alt](images/{sha}.png)\n<!-- vc:gorsel {sha} -->\n*Görsel:* alt")
    check("block without text", images.format_block(sha, "jpg", "", "") == f"![](images/{sha}.jpg)\n<!-- vc:gorsel {sha} -->")


def test_tokenize_and_resolve() -> None:
    png = noise_png(90, 90, seed=9)
    uri = "data:image/png;base64," + base64.b64encode(png).decode()
    md = f"<!-- Slide number: 1 -->\n\n![Picture 3]({uri})\n\n<!-- Slide number: 2 -->\n![Şekil: ağ]({uri})\n![broken](data:image/png;base64,AAAA)\ntext"
    col = images.Collector(enabled=True)
    tok = images.tokenize_data_uri_images(md, col)
    check("base64 is gone from the text", "base64" not in tok and tok.count("<!-- vc:img ") == 3)
    check("pages come from slide markers", [c.page for c in col.candidates] == [1, 2, 2])
    with tempfile.TemporaryDirectory() as td:
        out, stats, names = images.resolve_tokens(tok, col.candidates, Path(td) / "images")
        sha = images.sha16(images.reencode(png)[0])
        check("files written once under images/", names == [f"{sha}.png"] and (Path(td) / "images" / names[0]).read_bytes()[:4] == b"\x89PNG")
        check(
            "auto alt dropped, author alt becomes Görsel line",
            f"![](images/{sha}.png)\n<!-- vc:gorsel {sha} -->\n\n" in out
            and f"![Şekil: ağ](images/{sha}.png)\n<!-- vc:gorsel {sha} -->\n*Görsel:* Şekil: ağ" in out,
        )
        check("stats count every filter", stats["found"] == 3 and stats["kept"] == 2 and stats["unreadable"] == 1)
        check("no token left", "vc:img" not in out)
    # an inline picture (DOCX equation image inside a sentence) stays on its line
    col2 = images.Collector(enabled=True)
    tok2 = images.tokenize_data_uri_images(f"hata ({'![](' + uri + ')'}) olarak", col2)
    with tempfile.TemporaryDirectory() as td:
        out2, _, _ = images.resolve_tokens(tok2, col2.candidates, Path(td) / "images")
        check("inline picture does not split the sentence", out2.count("\n") == 0 and out2.startswith("hata (![](images/") and out2.endswith(") olarak"))


def test_pptx_end_to_end() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    import convert

    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        logo = td / "logo.png"
        logo.write_bytes(noise_png(100, 100, seed=11))
        fig = td / "fig.png"
        fig.write_bytes(noise_png(300, 200, seed=12))
        prs = Presentation()
        for i in range(4):
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = f"Slayt {i + 1}"
            logo_shape = slide.shapes.add_picture(str(logo), Inches(0.2), Inches(0.2), Inches(0.5), Inches(0.5))
            logo_shape._element.nvPicPr.cNvPr.set("descr", "logo içeren bir resim Yapay zeka tarafından oluşturulan içerik yanlış olabilir.")
            if i == 1:
                pic = slide.shapes.add_picture(str(fig), Inches(1), Inches(2), Inches(6), Inches(4))
                pic._element.nvPicPr.cNvPr.set("descr", "Gradyan inişin yol haritası")
        path = td / "deck.pptx"
        prs.save(str(path))

        images.COLLECTOR.reset(True)
        text = convert.convert_with_markitdown(path)
        check("pptx: 5 pictures collected, no base64 in text", len(images.COLLECTOR.candidates) == 5 and "base64" not in text)
        out, stats, names = images.resolve_tokens(text, images.COLLECTOR.candidates, td / "images", total_pages=4)
        check("pptx: logo on 4 slides dropped as template, figure kept", stats["template"] == 4 and stats["kept"] == 1 and len(names) == 1)
        check(
            "pptx: figure sits on slide 2 with author alt text",
            "Gradyan inişin yol haritası" in out and out.index("Slide number: 2") < out.index("vc:gorsel") < out.index("Slide number: 3"),
        )
        check("pptx: no dangling Name.jpg link", ".jpg)" not in out and "Resim" not in out)
        images.COLLECTOR.reset(False)


def test_docling_tokens() -> None:
    try:
        from docling_core.types.doc import BoundingBox, DocItemLabel, DoclingDocument, ImageRef, ProvenanceItem, Size
        from PIL import Image
    except ImportError:
        print("skip docling tests (docling_core missing)")
        return
    doc = DoclingDocument(name="t")
    doc.add_page(page_no=1, size=Size(width=600, height=800))
    doc.add_page(page_no=2, size=Size(width=600, height=800))
    doc.add_heading("Giriş", level=1)
    doc.add_text(label=DocItemLabel.TEXT, text="Önce metin.")
    for page, seed in ((1, 21), (2, 22)):
        pil = Image.open(io.BytesIO(noise_png(90, 70, seed=seed)))
        pic = doc.add_picture(
            image=ImageRef.from_pil(pil, dpi=72),
            prov=ProvenanceItem(page_no=page, bbox=BoundingBox(l=0, t=100, r=300, b=0), charspan=(0, 0)),
        )
        cap = doc.add_text(label=DocItemLabel.CAPTION, text=f"Şekil {page}: bir açıklama. İkinci cümle.")
        pic.captions.append(cap.get_ref())
    doc.add_text(label=DocItemLabel.TEXT, text="Sonra metin.")
    stock = doc.export_to_markdown()
    col = images.Collector(enabled=True)
    tok = images.docling_markdown_with_tokens(doc, col)
    check("docling: one token per picture, in order", tok.count("<!-- vc:img ") == 2 and tok.index("vc:img 0") < tok.index("vc:img 1"))
    restored = images.TOKEN_RE.sub("<!-- image -->", tok)
    check("docling: output equals export_to_markdown apart from the pictures", restored == stock)
    check(
        "docling: page, caption and area come from the picture itself",
        [c.page for c in col.candidates] == [1, 2]
        and col.candidates[0].alt == "Şekil 1: bir açıklama."
        and col.candidates[0].alt_in_text
        and col.candidates[0].area_frac is not None
        and abs(col.candidates[0].area_frac - 0.0625) < 1e-6,
    )
    with tempfile.TemporaryDirectory() as td:
        out, stats, names = images.resolve_tokens(tok, col.candidates, Path(td) / "images")
        check("docling: caption is alt only, no repeated Görsel line", "![Şekil 1: bir açıklama.](images/" in out and "*Görsel:*" not in out and len(names) == 2)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    test_alt_filter()
    test_reencode()
    test_equation_images()
    test_select_filters()
    test_format_block()
    test_tokenize_and_resolve()
    test_pptx_end_to_end()
    test_docling_tokens()
    print(f"\n{len(failures)} başarısız" if failures else "\nhepsi geçti")
    sys.exit(1 if failures else 0)
